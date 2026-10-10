from io import BytesIO
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.models.export import MaterialsExport

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_filename(data: MaterialsExport) -> str:
    def sanitize(value: str | None, fallback: str) -> str:
        ascii_value = unicodedata.normalize("NFKD", value or "").encode("ascii", "ignore").decode()
        return re.sub(r"[^a-zA-Z0-9_-]+", "_", ascii_value).strip("_-")[:60] or fallback
    return f"colcams_{sanitize(data.client, 'cliente')}_{sanitize(data.project or data.site, 'proyecto')}_{data.date.isoformat()}.xlsx"


def append_row(sheet, values):
    # All user/BOM text is literal, never an Excel formula. Strip illegal XML controls.
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value)[:32767] if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def style_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def generate_materials_excel(data: MaterialsExport) -> bytes:
    """Render the current estimate snapshot; do not invoke the quote engine."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    append_row(summary, ["Concepto", "Valor"])
    estimate = data.estimate
    rows = [
        ("Cliente", data.client), ("Proyecto", data.project), ("Dirección / sitio", data.site),
        ("Fecha", data.date), ("Cantidad total de cámaras", len(data.cameras)),
        ("Cámaras interiores", sum(c.environment == "indoor" for c in data.cameras)),
        ("Cámaras exteriores", sum(c.environment == "outdoor" for c in data.cameras)),
        ("Días de grabación", data.retention_days),
        ("Modo de grabación", {"continuous": "Continua", "events": "Solo ante eventos", "undefined": "A definir"}[data.recordingMode]),
        ("Horas por día usadas para almacenamiento", estimate.storage_hours_per_day if estimate.storage_hours_per_day is not None else data.recording_hours_per_day),
        ("NVR recomendado (canales)", estimate.nvr_channels),
        ("Almacenamiento estimado (TB)", estimate.storage_tb_raw),
        ("Almacenamiento seleccionado (TB)", estimate.storage_tb_selected),
        ("Switch PoE recomendado (puertos)", estimate.poe_switch_ports_selected),
        ("Metros estimados de cable", estimate.estimated_cable_m),
        ("Mano de obra", estimate.labor_cost), ("Costo estimado", estimate.total_cost),
        ("Margen", estimate.margin_percent / 100), ("Precio de venta", estimate.sale_price),
    ]
    alarm = estimate.alarm
    if alarm is not None:
        config = alarm.configuration
        system_labels = {"wired": "Cableado", "wireless": "Inalámbrico", "hybrid": "Híbrido"}
        communication_labels = {"ethernet": "Ethernet / IP", "wifi": "Wi-Fi", "lte": "LTE / 4G", "telephone": "Línea telefónica"}
        rows.extend([
            ("Sistema de alarma", system_labels[alarm.systemType]),
            ("Dispositivos de alarma", alarm.deviceCount),
            ("Zonas requeridas", alarm.zonesRequired),
            ("Reserva de ampliación de alarma (%)", alarm.expansionReservePercent),
            ("Zonas con reserva", alarm.zonesWithReserve),
            ("Panel recomendado (zonas)", alarm.recommendedPanelZones if alarm.recommendedPanelZones is not None else "Superior / revisión manual"),
            ("Panel seleccionado (zonas)", alarm.selectedPanelZones if alarm.selectedPanelZones is not None else "Revisión manual"),
            ("Panel modificado manualmente", "Sí" if alarm.panelOverridden else "No"),
            ("Expansores de alarma", alarm.expanderCount),
            ("Teclados de alarma", config.keypadCount),
            ("Sirenas interiores", config.indoorSirens),
            ("Sirenas exteriores", config.outdoorSirens),
            ("Comunicación de alarma", " + ".join(communication_labels[c] for c in config.communications) or "Sin comunicación remota"),
            ("Particiones de alarma", config.partitions),
            ("Autonomía de alarma (h)", config.backupAutonomyHours),
            ("Carga estimada de alarma (W)", alarm.estimatedLoadW),
            ("Batería aproximada (Ah, 12 V)", alarm.batteryAhApprox),
            ("Batería seleccionada (Ah, 12 V)", alarm.batteryAhSelected if alarm.batteryAhSelected is not None else "Superior / revisión manual"),
            ("Fuente auxiliar de alarma", "Sí" if alarm.auxiliaryPowerRequired else "No"),
            ("Supervisión anti-sabotaje", "Sí" if config.tamperRequired else "No"),
            ("Cable estimado de alarma (m)", alarm.estimatedCableM),
        ])
    for label, value in rows:
        if value is None or value == "":
            continue
        append_row(summary, [label, value])
        cell = summary.cell(summary.max_row, 2)
        if label in {"Mano de obra", "Costo estimado", "Precio de venta"}:
            cell.number_format = '"$" #,##0.00'
        elif label == "Margen":
            cell.number_format = "0.00%"
        elif label == "Fecha":
            cell.number_format = "yyyy-mm-dd"
    for warning in estimate.warnings:
        append_row(summary, ["Advertencia del resultado", warning])
    for pending in estimate.technical_pending:
        append_row(summary, ["Pendiente técnico / comercial", pending])
    if alarm is not None:
        for warning in alarm.warnings:
            append_row(summary, ["Advertencia de alarma", warning])
    append_row(summary, ["Alcance", "Estimación preliminar. Los costos no incluyen precios de catálogo. Las recomendaciones de lente requieren revisión técnica."])
    materials = workbook.create_sheet("Materiales")
    append_row(materials, ["Categoría", "Descripción", "Cantidad", "Unidad", "Observaciones"])
    categories = {
        "camera": "Cámaras", "nvr": "Grabación", "storage": "Almacenamiento", "switch": "Red", "cable": "Cableado CCTV",
        "alarm_panel": "Panel de alarma", "alarm_expander": "Expansión de alarma", "alarm_keypad": "Teclados de alarma",
        "alarm_sensor": "Sensores de alarma", "alarm_siren": "Sirenas de alarma", "alarm_communication": "Comunicación de alarma",
        "alarm_power": "Alimentación de alarma", "alarm_battery": "Respaldo de alarma", "alarm_accessories": "Accesorios de alarma", "alarm_cable": "Cableado de alarma",
    }
    for item in estimate.bom:
        # Unit metadata is absent in the current BOM; only map known categories.
        unit = item.unit or ("m" if item.category in {"cable", "alarm_cable"} else "un" if item.category in categories else None)
        append_row(materials, [categories.get(item.category, item.category), item.description, item.quantity, unit, item.observations or ("Estimado" if item.source == "engine" else None)])
    cameras = workbook.create_sheet("Cámaras")
    append_row(cameras, ["ID", "Nombre", "Ubicación", "Ambiente", "Formato", "Resolución (MP)", "Conectividad", "Distancia de cableado (m)", "Alcance visual", "Objetivo de imagen", "Distancia al objetivo (m)", "Objetivo nocturno", "Iluminación nocturna", "Color nocturno", "Evento", "Objetivo del evento", "Acciones", "Pendientes", "Recomendación preliminar de lente", "Notas"])
    ranges = {"near": "Vista general / corta distancia", "medium": "Distancia media", "far": "Objetivo lejano", "mixed": "Cercano y lejano"}
    formats = {"turret": "Torreta", "bullet": "Tipo bala", "dome": "Domo", "other": "Otro"}
    objectives = {"undefined": "Sin definir", "overview": "Vista general", "recognize": "Reconocer", "identify": "Identificar"}
    yes_no = {"yes": "Sí", "no": "No", "undefined": "A definir"}
    lighting = {"none": "Sin iluminación", "permanent": "Iluminación permanente", "motion": "Iluminación que se enciende por movimiento", "unknown": "A verificar"}
    events = {"none": "Sin detección adicional solicitada", "motion": "Movimiento dentro de la imagen", "line_crossing": "Cruce de una línea virtual", "intrusion_zone": "Ingreso o permanencia en una zona definida", "undefined": "A definir"}
    targets = {"any": "Cualquier movimiento u objeto", "person": "Personas", "vehicle": "Vehículos", "person_vehicle": "Personas y vehículos", "undefined": "A definir"}
    actions = {"mobile_notification": "Aviso al celular", "external_siren": "Activación de una sirena externa"}
    for camera in data.cameras:
        append_row(cameras, [
            camera.id, camera.name, camera.location,
            "Interior" if camera.environment == "indoor" else "Exterior", formats[camera.formFactor],
            camera.resolutionMp, "PoE" if camera.connectivity == "poe" else "Wi-Fi", camera.distanceM,
            ranges[camera.viewingRange], objectives[camera.imageObjective], camera.targetDistanceM,
            yes_no[camera.nightObjectiveRequired], lighting[camera.nightLighting],
            "No, acepta blanco y negro" if camera.nightColorRequired == "no" else yes_no[camera.nightColorRequired],
            events[camera.detectionEvent], targets[camera.detectionTarget],
            "; ".join(actions[action] for action in camera.eventActions),
            "\n".join(camera.technicalPending), camera.lensRecommendation, camera.notes,
        ])
    style_sheet(summary, [40, 85])
    style_sheet(materials, [22, 55, 15, 12, 35])
    style_sheet(cameras, [12, 25, 30, 16, 16, 18, 18, 24, 35, 24, 30, 24, 35, 24, 45, 30, 45, 85, 85, 45])
    if alarm is not None:
        alarm_sheet = workbook.create_sheet("Alarma")
        append_row(alarm_sheet, ["Elemento", "Ubicación", "Tipo", "Conexión", "Cantidad", "Zona", "Observaciones"])
        device_labels = {"pir_indoor": "Sensor PIR interior", "pir_outdoor": "Sensor PIR exterior", "magnetic_contact": "Contacto magnético", "beam": "Barrera infrarroja", "glass_break": "Detector de rotura de vidrio", "smoke": "Detector de humo", "gas": "Detector de gas", "flood": "Detector de inundación", "panic_button": "Pulsador de pánico", "other": "Otro"}
        for device in alarm.devices:
            append_row(alarm_sheet, [device.name, device.location, device_labels[device.type], "Cableado" if device.connection == "wired" else "Inalámbrico", device.quantity, device.zoneLabel, device.notes])
        for warning in alarm.warnings:
            append_row(alarm_sheet, ["Advertencia técnica", None, None, None, None, None, warning])
        style_sheet(alarm_sheet, [30, 30, 35, 20, 15, 30, 90])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
