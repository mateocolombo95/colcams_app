import asyncio
from datetime import date
from io import BytesIO
import json

from openpyxl import load_workbook

from app.main import app
from app.models.export import MaterialsExport
from app.models.quote import CCTVRequirements
from app.services.quote_engine import estimate_quote
from app.services.excel_export import export_filename, generate_materials_excel


def snapshot():
    estimate = estimate_quote(CCTVRequirements(camera_count=1))
    return MaterialsExport(
        client="José / Cliente", site="Portón: principal", date=date(2026, 10, 6),
        estimate=estimate,
        cameras=[dict(id="C1", name="Portón", environment="outdoor", formFactor="bullet",
                      resolutionMp=4, connectivity="poe", distanceM=32, viewingRange="far",
                      targetDistanceM=28, lensRecommendation="Considerar varifocal 6–12 mm.")],
    )


def test_workbook_uses_current_bom_and_camera_details():
    data = snapshot()
    # A changed BOM snapshot must be exported as-is, not recomputed by the exporter.
    data.estimate.bom[0].quantity = 17
    book = load_workbook(BytesIO(generate_materials_excel(data)))
    assert book.sheetnames == ["Resumen", "Materiales", "Cámaras"]
    summary = dict(book["Resumen"].values)
    assert summary["Cantidad total de cámaras"] == 1
    assert summary["Cámaras exteriores"] == 1
    assert "Proyecto" not in summary
    assert "Días de grabación" not in summary
    assert book["Materiales"].max_row == len(data.estimate.bom) + 1
    for row, item in zip(list(book["Materiales"].values)[1:], data.estimate.bom):
        assert row[1] == item.description
        assert row[2] == item.quantity
    camera = dict(zip(next(book["Cámaras"].values), list(book["Cámaras"].values)[1]))
    assert camera["Distancia de cableado (m)"] == 32
    assert camera["Distancia al objetivo (m)"] == 28
    assert camera["Recomendación preliminar de lente"] == data.cameras[0].lensRecommendation
    assert book["Materiales"].freeze_panes == "A2"
    assert book["Materiales"].auto_filter.ref
    assert book["Resumen"]["B2"].data_type == "s"


def test_missing_details_formula_text_and_safe_filename():
    data = snapshot()
    data.cameras[0].targetDistanceM = None
    data.cameras[0].name = '=HYPERLINK("https://example.com")\x01'
    data.client = "=SUM(1,2)"
    book = load_workbook(BytesIO(generate_materials_excel(data)))
    camera = dict(zip(next(book["Cámaras"].values), list(book["Cámaras"].values)[1]))
    assert camera["Distancia al objetivo (m)"] is None
    assert book["Cámaras"]["B2"].data_type == "s"
    assert "\x01" not in book["Cámaras"]["B2"].value
    assert book["Resumen"]["B2"].data_type == "s"
    data = snapshot()
    assert export_filename(data) == "colcams_Jose_Cliente_Porton_principal_2026-10-06.xlsx"
    data.cameras = []
    data.estimate.bom = []
    assert load_workbook(BytesIO(generate_materials_excel(data)))["Materiales"].max_row == 1


def test_endpoint_returns_valid_xlsx_and_validation_errors():
    async def request(payload):
        sent = []
        async def receive():
            return {"type": "http.request", "body": json.dumps(payload).encode(), "more_body": False}
        async def send(message):
            sent.append(message)
        await app({"type": "http", "asgi": {"version": "3.0"}, "http_version": "1.1",
                   "method": "POST", "scheme": "http", "path": "/api/v1/exports/materials.xlsx",
                   "raw_path": b"/api/v1/exports/materials.xlsx", "query_string": b"",
                   "headers": [(b"content-type", b"application/json")],
                   "client": ("127.0.0.1", 1), "server": ("test", 80)}, receive, send)
        return sent
    messages = asyncio.run(request(snapshot().model_dump(mode="json")))
    assert messages[0]["status"] == 200
    headers = dict(messages[0]["headers"])
    assert b"spreadsheetml.sheet" in headers[b"content-type"]
    assert b".xlsx" in headers[b"content-disposition"]
    assert load_workbook(BytesIO(messages[1]["body"])).sheetnames == ["Resumen", "Materiales", "Cámaras"]
    assert asyncio.run(request({}))[0]["status"] == 422


def test_new_camera_requirements_pending_and_storage_assumptions_are_exported():
    data = snapshot()
    data.recordingMode = "events"
    data.estimate = estimate_quote(CCTVRequirements(camera_count=1, recordingMode="events", recording_hours_per_day=4))
    camera = data.cameras[0]
    camera.imageObjective = "identify"
    camera.nightObjectiveRequired = "yes"
    camera.nightLighting = "permanent"
    camera.nightColorRequired = "no"
    camera.detectionEvent = "intrusion_zone"
    camera.detectionTarget = "vehicle"
    camera.eventActions = ["mobile_notification", "external_siren"]
    camera.technicalPending = ["Verificar integración de sirena externa."]
    book = load_workbook(BytesIO(generate_materials_excel(data)))
    details = dict(zip(next(book["Cámaras"].values), list(book["Cámaras"].values)[1]))
    assert details["Objetivo de imagen"] == "Identificar"
    assert details["Alcance visual"] == "Objetivo lejano"
    assert details["Distancia al objetivo (m)"] == 28
    assert details["Objetivo nocturno"] == "Sí"
    assert details["Iluminación nocturna"] == "Iluminación permanente"
    assert details["Color nocturno"] == "No, acepta blanco y negro"
    assert details["Evento"] == "Ingreso o permanencia en una zona definida"
    assert details["Objetivo del evento"] == "Vehículos"
    assert "Aviso al celular" in details["Acciones"]
    assert "sirena externa" in details["Acciones"]
    assert details["Pendientes"] == camera.technicalPending[0]
    summary = dict(book["Resumen"].values)
    assert summary["Modo de grabación"] == "Solo ante eventos"
    assert summary["Horas por día usadas para almacenamiento"] == 24
    assert any("grabación continua como referencia" in row[1] for row in book["Resumen"].values if isinstance(row[1], str))


def test_export_route_normalizes_camera_pending_without_recalculating_bom(monkeypatch):
    from app.api.routes.exports import export_materials
    from app.services import quote_engine
    data = snapshot()
    data.cameras[0].imageObjective = "identify"
    data.cameras[0].targetDistanceM = None
    data.cameras[0].eventActions = ["external_siren"]
    data.estimate.bom[0].quantity = 17
    def unexpected_recalculation(*args, **kwargs):
        raise AssertionError("El exportador no debe recalcular la estimación.")
    monkeypatch.setattr(quote_engine, "estimate_quote", unexpected_recalculation)
    response = export_materials(data)
    book = load_workbook(BytesIO(response.body))
    details = dict(zip(next(book["Cámaras"].values), list(book["Cámaras"].values)[1]))
    assert "distancia de identificación" in details["Pendientes"]
    assert "compatibilidad" in details["Pendientes"]
    assert list(book["Materiales"].values)[1][2] == 17
