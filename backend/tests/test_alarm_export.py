from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.models.alarm import AlarmConfiguration, AlarmDeviceRequirement
from app.models.export import MaterialsExport
from app.models.quote import CCTVRequirements
from app.services.excel_export import generate_materials_excel
from app.services.quote_engine import estimate_quote


def exported_alarm_snapshot(enabled=True):
    alarm = AlarmConfiguration(
        enabled=enabled, systemType="hybrid", panelMode="manual", panelZones=8,
        communications=["ethernet", "lte"], partitions=2, expansionReservePercent=20,
        devices=[
            AlarmDeviceRequirement(id="A1", name="Hall", location="Planta baja", type="pir_indoor", quantity=4, connection="wired"),
            AlarmDeviceRequirement(id="A2", name="Accesos", type="magnetic_contact", quantity=5, connection="wireless"),
            AlarmDeviceRequirement(id="A3", name="Pánico", type="panic_button", quantity=1, connection="wired"),
        ],
    )
    estimate = estimate_quote(CCTVRequirements(camera_count=2, alarm=alarm))
    return MaterialsExport(date=date(2026, 10, 6), estimate=estimate, cameras=[])


def test_alarm_export_reads_result_snapshot_and_unified_bom():
    snapshot = exported_alarm_snapshot()
    book = load_workbook(BytesIO(generate_materials_excel(snapshot)))
    assert book.sheetnames == ["Resumen", "Materiales", "Cámaras", "Alarma"]
    summary = dict(book["Resumen"].values)
    assert summary["Sistema de alarma"] == "Híbrido"
    assert summary["Zonas requeridas"] == 10
    assert summary["Panel seleccionado (zonas)"] == 8
    assert summary["Expansores de alarma"] == 1
    assert summary["Comunicación de alarma"] == "Ethernet / IP + LTE / 4G"
    details = list(book["Alarma"].values)
    assert details[1][:6] == ("Hall", "Planta baja", "Sensor PIR interior", "Cableado", 4, snapshot.estimate.alarm.devices[0].zoneLabel)
    assert details[2][3] == "Inalámbrico"
    materials = list(book["Materiales"].values)[1:]
    assert len(materials) == len(snapshot.estimate.bom)
    for row, item in zip(materials, snapshot.estimate.bom):
        assert row[1:3] == (item.description, item.quantity)
        if item.unit:
            assert row[3] == item.unit
    # Change the saved snapshot: exporter must NOT invoke the engine again.
    snapshot.estimate.alarm.zonesRequired = 99
    snapshot.estimate.bom[-1].quantity = 123
    changed = load_workbook(BytesIO(generate_materials_excel(snapshot)))
    assert dict(changed["Resumen"].values)["Zonas requeridas"] == 99
    assert list(changed["Materiales"].values)[-1][2] == 123


def test_disabled_alarm_exports_original_three_sheets_without_alarm_materials():
    snapshot = exported_alarm_snapshot(enabled=False)
    book = load_workbook(BytesIO(generate_materials_excel(snapshot)))
    assert book.sheetnames == ["Resumen", "Materiales", "Cámaras"]
    assert "Sistema de alarma" not in dict(book["Resumen"].values)
    assert not any("alarm" in item.category for item in snapshot.estimate.bom)
